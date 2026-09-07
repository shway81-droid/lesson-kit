import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from read_progress_plan import Lesson, merge_blocks, read_lessons

HEADER = ["출력순", "학년", "학기", "편제", "단원", "소단원", "학습내용",
          "쪽수", "보조쪽수", "해당차시", "전체차시", "준비물"]

ROWS = [
    [1, 5, 2, "사회", "1. 유적과 유물", "단원 도입", "단원 도입", "8~11", None, 1, 20, "교과서"],
    [2, 5, 2, "사회", "1. 유적과 유물", "① 선사 시대", "구석기 시대", "12~14", None, 2, 20, "교과서"],
    # 2차시짜리 수업. 진도표는 같은 쪽수로 두 줄에 적는다.
    [3, 5, 2, "사회", "1. 유적과 유물", "② 고대", "통일신라", "40~45", None, 10, 20, "교과서"],
    [4, 5, 2, "사회", "1. 유적과 유물", "② 고대", "통일신라", "40~45", None, 11, 20, "교과서"],
    [5, 5, 2, "사회", "2. 달라지는 시대", "창의 융합", "창의 융합", "114~115", None, 13, 14, "교과서"],
]


@pytest.fixture
def 진도표(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "진도표.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    wb.save(path)
    return path


def test_차시를_모두_읽는다(진도표):
    lessons = read_lessons(진도표)
    assert len(lessons) == 5
    assert lessons[1].title == "구석기 시대"
    assert (lessons[1].page_from, lessons[1].page_to) == (12, 14)
    assert lessons[1].lesson_no == 2
    assert lessons[1].unit_total == 20


def test_열_순서가_아니라_이름으로_찾는다(tmp_path):
    # 출판사가 열을 옮겨도 읽혀야 한다
    from openpyxl import Workbook

    path = tmp_path / "뒤바뀐.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["쪽수", "학습내용", "단원", "소단원", "해당차시", "전체차시"])
    ws.append(["12~14", "구석기 시대", "1. 유적과 유물", "① 선사 시대", 2, 20])
    wb.save(path)

    lessons = read_lessons(path)
    assert lessons[0].title == "구석기 시대"
    assert lessons[0].page_from == 12


def test_열이_없으면_멈춘다(tmp_path):
    # 쪽수 열이 없는데 그냥 진행하면 엉뚱한 범위를 잘라 내고도 정상 종료한다
    from openpyxl import Workbook

    path = tmp_path / "쪽수없음.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["단원", "소단원", "학습내용", "해당차시", "전체차시"])
    ws.append(["1. 유적과 유물", "① 선사 시대", "구석기 시대", 2, 20])
    wb.save(path)

    with pytest.raises(ValueError, match="쪽수"):
        read_lessons(path)


def test_한_쪽짜리_차시도_읽는다(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "한쪽.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append([1, 5, 2, "사회", "1. 단원", "소단원", "제목", "70", None, 19, 20, ""])
    wb.save(path)

    lessons = read_lessons(path)
    assert (lessons[0].page_from, lessons[0].page_to) == (70, 70)


def test_쪽수가_같은_연속_차시를_합친다(진도표):
    lessons = merge_blocks(read_lessons(진도표))
    # 통일신라 10·11차시가 하나로 합쳐져 5개 → 4개
    assert len(lessons) == 4
    통일신라 = [l for l in lessons if l.title == "통일신라"]
    assert len(통일신라) == 1
    assert 통일신라[0].lesson_no == 10


def test_단원이_다르면_쪽수가_같아도_합치지_않는다():
    lessons = [
        Lesson("1. 가", "소", "제목", 10, 12, 1, 20),
        Lesson("2. 나", "소", "제목", 10, 12, 1, 14),
    ]
    assert len(merge_blocks(lessons)) == 2


def test_pdf_범위는_offset_만큼_당긴다():
    lesson = Lesson("1. 유적과 유물", "② 고대", "고조선", 22, 25, 5, 20)
    assert lesson.pdf_range(1) == (21, 24)
    assert lesson.pdf_range(0) == (22, 25)


def test_slug_은_정렬되고_경로에_쓸_수_있다():
    lesson = Lesson("1. 유적과 유물", "② 고대", "고조선은 어떤 나라일까요", 22, 25, 5, 20)
    slug = lesson.slug
    assert slug.startswith("1단원_05차시_")
    assert "고조선은" in slug
    for bad in ' /\\:*?"<>|':
        assert bad not in slug
